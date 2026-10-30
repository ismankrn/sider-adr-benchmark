"""Build the Word manuscript for Molecular Informatics from main.tex (pandoc + python-docx post-processing)."""
import re, subprocess, sys, pathlib
HERE = pathlib.Path(__file__).resolve().parent
tex = (HERE / "main.tex").read_text()
tex = re.sub(r"^\\(newcommand\{\\(ecfp|repourl|repodoi|fittable|todo)\}|input\{tables/numbers\}|newsavebox|usepackage\[[^\]]*\]\{lineno\}|linenumbers|onehalfspacing|setlength\{\\emergencystretch\}).*$\n", "", tex, flags=re.M)
aux = (HERE / "main.aux").read_text()
# --- resolve \ref / \eqref numbers from the LaTeX aux file
labels = dict(re.findall(r"\\newlabel\{([^}]+)\}\{\{([^}]*)\}", aux))
def ref(m):
    key = m.group(2); num = labels.get(key, "??")
    num = re.sub(r"\\[a-zA-Z]+\s*", "", num).strip("{}")
    return f"({num})" if m.group(1) == "eqref" else num
tex = re.sub(r"\\(ref|eqref)\{([^}]+)\}", ref, tex)
tex = tex.replace("~", "\u00a0")  # non-breaking spaces in "Table 3", "Section 2"
# --- numbers macros
nums = dict(re.findall(r"\\newcommand\{\\(\w+)\}\{([^}]*)\}", (HERE / "tables/numbers.tex").read_text()))
for k, v in sorted(nums.items(), key=lambda kv: -len(kv[0])):
    tex = tex.replace("\\" + k + "{}", v); tex = re.sub(r"\\" + k + r"(?![A-Za-z])", v, tex)
tex = tex.replace(r"\ecfp{}", "ECFP4").replace(r"\ecfp", "ECFP4")
tex = tex.replace(r"\repourl{}", r"\url{https://github.com/ismankrn/sider-adr-benchmark}")
tex = tex.replace(r"\repodoi{}", r"\href{https://doi.org/10.5281/zenodo.23798100}{10.5281/zenodo.23798100}")
# --- inline the generated tables, drop size/fit wrappers
def inp(m):
    return (HERE / f"tables/{m.group(1)}.tex").read_text()
tex = re.sub(r"\\input\{tables/(\w+)\}", inp, tex)
tex = re.sub(r"\\(small|scriptsize|footnotesize|centering)\s*", "", tex)
tex = re.sub(r"\\fittable\{", "{", tex)
tex = re.sub(r"\\setlength\{\\tabcolsep\}\{[^}]*\}|\\renewcommand\{\\arraystretch\}\{[^}]*\}", "", tex)
tex = re.sub(r"\\cmidrule\(lr\)\{[^}]*\}", "", tex)
# --- figures: use PNGs
tex = re.sub(r"\\includegraphics\[([^\]]*)\]\{([^}]+)\}", lambda m: f"\\includegraphics[{m.group(1)}]{{figures/{pathlib.Path(m.group(2)).name}.png}}" if not m.group(2).startswith("rank") else f"\\includegraphics[{m.group(1)}]{{../code/figures/attribution/{m.group(2)}.png}}", tex)
tex = re.sub(r"\\graphicspath\{[^\n]*\}\n", "", tex)
# --- drop LaTeX-only preamble lines
for pat in [r"\\newcommand\{\\(ecfp|repourl|repodoi)\}[^\n]*\n", r"\\input\{tables/numbers\}\n", r"\\newsavebox[^\n]*\n", r"\\usepackage\[[^\]]*\]\{lineno\}\n", r"\\linenumbers\n", r"\\onehalfspacing\n", r"\\setlength\{\\emergencystretch\}\{1em\}\n", r"\\newsavebox\{\\tablebox\}\n", r"\\newcommand\{\\fittable\}[^\n]*\n", r"\\newcommand\{\\todo\}[^\n]*\n"]:
    tex = re.sub(pat, "", tex)
(HERE / "build" ).mkdir(exist_ok=True)
(HERE / "build/main_docx.tex").write_text(tex)
csl = sys.argv[1] if len(sys.argv) > 1 else "angewandte-chemie.csl"
cmd = ["pandoc", "build/main_docx.tex", "-f", "latex", "-t", "docx", "--citeproc", "--bibliography", "references.bib", "--csl", csl,
       "--resource-path", ".:figures:../code/figures/attribution", "-o", "build/manuscript_MolInf.docx"]
if (HERE / "reference.docx").exists(): cmd += ["--reference-doc", "reference.docx"]
print(" ".join(cmd)); subprocess.run(cmd, cwd=HERE, check=True)
# --- post-processing with python-docx
import docx
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
d = docx.Document(HERE / "build/manuscript_MolInf.docx")
for s in d.styles:
    try:
        if s.type == 1: s.font.name = "Arial"
    except Exception: pass
d.styles["Normal"].font.size = Pt(10)
for p in d.paragraphs:
    for r in p.runs:
        if r._element.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for shp in r._element.xpath(".//wp:extent"):
                cx = int(shp.get("cx")); cy = int(shp.get("cy"))
                maxw = Cm(17.5)
                if cx > maxw:
                    shp.set("cy", str(int(cy * maxw / cx))); shp.set("cx", str(int(maxw)))
                    for ext in r._element.xpath(".//a:ext"):
                        ext.set("cx", str(int(maxw))); ext.set("cy", str(int(cy * maxw / cx)))
for t in d.tables:
    t.style = d.styles["Table Grid"] if "Table Grid" in [s.name for s in d.styles] else t.style
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs: r.font.size = Pt(8)
# --- title page: rebuild the author block with superscript footnote letters and one line per affiliation
import copy
for p in d.paragraphs:
    if p.style.name == "Author":
        for r in list(p.runs): r._element.getparent().remove(r._element)
        def sup(par, txt): rr = par.add_run(txt); rr.font.superscript = True
        p.add_run("Isman Kurniawan"); sup(p, "[a],*"); p.add_run(", Angel Metanosa Afinda"); sup(p, "[b]"); p.add_run(", Aaz Muhammad Hafidz Azis"); sup(p, "[a]")
        aff = ["[a] School of Computing, Telkom University, Jl. Telekomunikasi No. 1, Bandung 40257, Indonesia",
               "[b] School of Electrical Engineering, Telkom University, Jl. Telekomunikasi No. 1, Bandung 40257, Indonesia",
               "* Corresponding author. E-mail: ismankrn@telkomuniversity.ac.id"]
        anchor = p
        for a in aff:
            np_ = copy.deepcopy(p._element); anchor._element.addnext(np_)
            from docx.text.paragraph import Paragraph
            q = Paragraph(np_, p._parent)
            for r in list(q.runs): r._element.getparent().remove(r._element)
            rr = q.add_run(a); rr.font.size = Pt(9); anchor = q
        break
# --- number the captions ("Table n." / "Figure n.") in order of appearance, label in bold
tn = fn = 0
for p in d.paragraphs:
    if p.style.name == "Table Caption":
        tn += 1; lab = f"Table {tn}. "
    elif p.style.name == "Image Caption":
        fn += 1; lab = f"Figure {fn}. "
    else:
        continue
    r = p.runs[0]._element if p.runs else None
    nr = p.add_run(lab); nr.bold = True
    if r is not None: r.addprevious(nr._element)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
# --- images placed inside table cells (multi-panel figure): limit each to one third of the text width
for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            for ext in cell._element.xpath(".//wp:extent"):
                cx = int(ext.get("cx")); cy = int(ext.get("cy")); maxw = Cm(5.6)
                if cx > maxw:
                    ext.set("cx", str(int(maxw))); ext.set("cy", str(int(cy * maxw / cx)))
            for ext in cell._element.xpath(".//a:ext"):
                cx = int(ext.get("cx")); cy = int(ext.get("cy")); maxw = Cm(5.6)
                if cx > maxw:
                    ext.set("cx", str(int(maxw))); ext.set("cy", str(int(cy * maxw / cx)))
d.save(HERE / "build/manuscript_MolInf.docx")
print(f"wrote build/manuscript_MolInf.docx with {tn} tables and {fn} figures numbered")
